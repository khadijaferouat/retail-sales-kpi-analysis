"""
Construit le classeur Excel d'analyse (données, tableau de bord KPI,
détection d'anomalie). Les tableaux croisés dynamiques sont ajoutés
ensuite par scripts/add_pivot_tables.py (LibreOffice / UNO).

Usage : python scripts/build_workbook.py
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_csv(ROOT / "data" / "ventes_retail_2025.csv", sep=";", decimal=",")
df["Date"] = pd.to_datetime(df["Date"])
N = len(df) + 1  # dernière ligne de données
MAGS = list(df["Magasin"].unique())
MOIS = sorted(df["Mois"].unique())

F = "Arial"
NAVY, TEAL, GREY, LIGHT = "1F3A5F", "2A9D8F", "6B7280", "EEF2F7"
f_title = Font(name=F, size=18, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color=GREY)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_b = Font(name=F, size=10)
f_bb = Font(name=F, size=10, bold=True)
f_sec = Font(name=F, size=12, bold=True, color=NAVY)
f_input = Font(name=F, size=10, bold=True, color="0000FF")
fill_h = PatternFill("solid", fgColor=NAVY)
fill_l = PatternFill("solid", fgColor=LIGHT)
fill_y = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="D1D5DB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0 €;-#,##0 €;"-"'
EUR2 = '#,##0.00 €;-#,##0.00 €;"-"'
INT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
D = "'Données'"

def header(ws, row, col, labels, width=None):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.border = f_h, fill_h, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_range(ws, r1, r2, c1, c2, fmt=None, font=f_b):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.font = box, font
            if fmt and c > c1:
                cell.number_format = fmt

wb = Workbook()

# ---------------------------------------------------------------- Lisez-moi
ws0 = wb.active
ws0.title = "Lisez-moi"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 3
ws0.column_dimensions["B"].width = 26
ws0.column_dimensions["C"].width = 95
ws0["B2"] = "Analyse des ventes retail 2025"
ws0["B2"].font = f_title
ws0["B3"] = "Suivi d'indicateurs (KPI) et détection d'anomalie par tableau croisé dynamique"
ws0["B3"].font = f_sub
lines = [
    ("Contexte", "Réseau fictif de 4 magasins (Paris, Lyon, Lille, Bordeaux), 5 catégories de produits, ventes journalières de l'année 2025 (7 300 lignes). Les données sont synthétiques et générées par scripts/generate_data.py."),
    ("Objectif", "Suivre la performance commerciale du réseau avec un tableau de bord d'indicateurs et repérer automatiquement les écarts anormaux entre magasins."),
    ("Résultat", "Le tableau croisé dynamique Magasin x Mois fait apparaître un décrochage de Lyon Part-Dieu en octobre 2025 : environ 15 % de chiffre d'affaires en moins par rapport à sa référence habituelle face au reste du réseau. Le panier moyen reste stable, la baisse vient donc de la fréquentation (nombre de transactions)."),
    ("", ""),
    ("Onglet Données", "Table brute filtrable. La colonne Marge est calculée (CA HT moins coût d'achat)."),
    ("Onglet Tableau de bord", "KPI globaux, évolution mensuelle et répartition par magasin. La cellule jaune C4 permet de filtrer tout le tableau de bord sur un magasin."),
    ("Onglet TCD", "Tableaux croisés dynamiques Excel (CA, transactions, part du CA par magasin). Clic droit puis Actualiser si les données changent."),
    ("Onglet Détection anomalie", "Méthode : pour chaque magasin et chaque mois, ratio CA magasin / CA du reste du réseau, comparé à la moyenne de ce ratio sur les 11 autres mois. Un écart inférieur au seuil (cellule jaune, -10 % par défaut) déclenche une alerte."),
    ("", ""),
    ("Légende", "Texte bleu sur fond jaune : paramètre modifiable. Texte noir : formule. Rouge : alerte."),
]
for i, (k, v) in enumerate(lines):
    ws0.cell(row=5 + i, column=2, value=k).font = f_bb
    c = ws0.cell(row=5 + i, column=3, value=v)
    c.font, c.alignment = f_b, Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[5 + i].height = 30 if len(v) > 95 else 16
ws0.row_dimensions[7].height = 45

# ---------------------------------------------------------------- Données
ws = wb.create_sheet("Données")
cols = list(df.columns) + ["Marge"]
header(ws, 1, 1, cols)
for r, row in enumerate(df.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v.to_pydatetime() if c == 1 else v)
    ws.cell(row=r, column=10, value=f"=H{r}-I{r}")
for r in range(2, N + 1):
    ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    for c in (8, 9, 10):
        ws.cell(row=r, column=c).number_format = "#,##0.00"
for c, w in zip("ABCDEFGHIJ", [12, 10, 17, 22, 14, 16, 11, 12, 12, 12]):
    ws.column_dimensions[c].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{N}"

rng = lambda col: f"{D}!${col}$2:${col}${N}"
R = {k: rng(c) for k, c in dict(mois="B", mag="C", cat="E", tx="F", qte="G",
                                   ca="H", cout="I", marge="J").items()}

# ---------------------------------------------------------------- Tableau de bord
db = wb.create_sheet("Tableau de bord")
db.sheet_view.showGridLines = False
db.column_dimensions["A"].width = 2
for c in range(2, 10):
    db.column_dimensions[L(c)].width = 18
db["B2"] = "Tableau de bord des ventes 2025"
db["B2"].font = f_title
db["B4"] = "Magasin :"
db["B4"].font = f_bb
db["C4"] = "Tous"
db["C4"].font, db["C4"].fill, db["C4"].border = f_input, fill_y, box
db["C4"].comment = Comment("Choisir un magasin dans la liste, ou Tous pour le réseau complet.", "Analyse")
dv = DataValidation(type="list", formula1='"Tous,' + ",".join(MAGS) + '"', allow_blank=False)
db.add_data_validation(dv)
dv.add("C4")
db["D4"] = '=IF($C$4="Tous","Réseau complet, 4 magasins","Filtre actif : "&$C$4)'
db["D4"].font = f_sub
SEL = 'IF($C$4="Tous","*",$C$4)'

kpis = [
    ("CA HT", f"=SUMIFS({R['ca']},{R['mag']},{SEL})", EUR),
    ("Marge brute", f"=SUMIFS({R['marge']},{R['mag']},{SEL})", EUR),
    ("Taux de marge", "=IFERROR(C7/B7,0)", PCT),
    ("Transactions", f"=SUMIFS({R['tx']},{R['mag']},{SEL})", INT),
    ("Panier moyen", "=IFERROR(B7/E7,0)", EUR2),
    ("Articles / ticket", f"=IFERROR(SUMIFS({R['qte']},{R['mag']},{SEL})/E7,0)", "0.00"),
]
for i, (lab, fml, fmt) in enumerate(kpis):
    c = 2 + i
    h = db.cell(row=6, column=c, value=lab)
    h.font, h.fill, h.alignment = Font(name=F, size=9, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=TEAL), Alignment(horizontal="center")
    v = db.cell(row=7, column=c, value=fml)
    v.font, v.fill, v.number_format = Font(name=F, size=14, bold=True, color=NAVY), fill_l, fmt
    v.alignment = Alignment(horizontal="center", vertical="center")
db.row_dimensions[7].height = 32

db["B10"] = "Évolution mensuelle"
db["B10"].font = f_sec
header(db, 11, 2, ["Mois", "CA HT", "Marge brute", "Taux de marge", "Transactions", "Panier moyen", "Var. CA M/M"])
r0 = 12
for i, m in enumerate(MOIS):
    r = r0 + i
    db.cell(row=r, column=2, value=m)
    db.cell(row=r, column=3, value=f"=SUMIFS({R['ca']},{R['mois']},$B{r},{R['mag']},{SEL})")
    db.cell(row=r, column=4, value=f"=SUMIFS({R['marge']},{R['mois']},$B{r},{R['mag']},{SEL})")
    db.cell(row=r, column=5, value=f"=IFERROR(D{r}/C{r},0)")
    db.cell(row=r, column=6, value=f"=SUMIFS({R['tx']},{R['mois']},$B{r},{R['mag']},{SEL})")
    db.cell(row=r, column=7, value=f"=IFERROR(C{r}/F{r},0)")
    db.cell(row=r, column=8, value="" if i == 0 else f"=IFERROR(C{r}/C{r-1}-1,0)")
rt = r0 + len(MOIS)
db.cell(row=rt, column=2, value="Total")
db.cell(row=rt, column=3, value=f"=SUM(C{r0}:C{rt-1})")
db.cell(row=rt, column=4, value=f"=SUM(D{r0}:D{rt-1})")
db.cell(row=rt, column=5, value=f"=IFERROR(D{rt}/C{rt},0)")
db.cell(row=rt, column=6, value=f"=SUM(F{r0}:F{rt-1})")
db.cell(row=rt, column=7, value=f"=IFERROR(C{rt}/F{rt},0)")
style_range(db, r0, rt, 2, 8)
for r in range(r0, rt + 1):
    for c, fmt in zip(range(3, 9), [EUR, EUR, PCT, INT, EUR2, PCT]):
        db.cell(row=r, column=c).number_format = fmt
for c in range(2, 9):
    db.cell(row=rt, column=c).font = f_bb
    db.cell(row=rt, column=c).fill = fill_l
db.conditional_formatting.add(f"H{r0+1}:H{rt-1}", CellIsRule(operator="lessThan", formula=["-0.05"], font=Font(name=F, color="C00000", bold=True)))

rm = rt + 3
db.cell(row=rm, column=2, value="CA HT par magasin et par mois (réseau complet)").font = f_sec
header(db, rm + 1, 2, ["Mois"] + MAGS + ["Total"])
for i, m in enumerate(MOIS):
    r = rm + 2 + i
    db.cell(row=r, column=2, value=m)
    for j, mag in enumerate(MAGS):
        db.cell(row=r, column=3 + j, value=f'=SUMIFS({R["ca"]},{R["mois"]},$B{r},{R["mag"]},{L(3+j)}${rm+1})')
    db.cell(row=r, column=3 + len(MAGS), value=f"=SUM(C{r}:{L(2+len(MAGS))}{r})")
rmt = rm + 2 + len(MOIS)
db.cell(row=rmt, column=2, value="Total")
for c in range(3, 4 + len(MAGS)):
    db.cell(row=rmt, column=c, value=f"=SUM({L(c)}{rm+2}:{L(c)}{rmt-1})")
style_range(db, rm + 2, rmt, 2, 3 + len(MAGS), EUR)
for c in range(2, 4 + len(MAGS)):
    db.cell(row=rmt, column=c).font, db.cell(row=rmt, column=c).fill = f_bb, fill_l

ch = LineChart()
ch.title, ch.y_axis.title, ch.height, ch.width = "CA HT mensuel", "€", 7.5, 16
ch.add_data(Reference(db, min_col=3, min_row=11, max_row=rt - 1), titles_from_data=True)
ch.set_categories(Reference(db, min_col=2, min_row=r0, max_row=rt - 1))
ch.legend = None
db.add_chart(ch, "J10")
bc = BarChart()
bc.title, bc.height, bc.width = "CA HT mensuel par magasin", 7.5, 16
bc.add_data(Reference(db, min_col=3, max_col=2 + len(MAGS), min_row=rm + 1, max_row=rmt - 1), titles_from_data=True)
bc.set_categories(Reference(db, min_col=2, min_row=rm + 2, max_row=rmt - 1))
db.add_chart(bc, f"J{rm}")

# ---------------------------------------------------------------- TCD (rempli par UNO)
tcd = wb.create_sheet("TCD")
tcd["A1"] = "Tableaux croisés dynamiques : source Données!A1:J" + str(N)
tcd["A1"].font = f_sec

# ---------------------------------------------------------------- Détection anomalie
an = wb.create_sheet("Détection anomalie")
an.sheet_view.showGridLines = False
an.column_dimensions["A"].width = 2
an.column_dimensions["B"].width = 30
for c in range(3, 9):
    an.column_dimensions[L(c)].width = 16
an["B2"] = "Détection d'anomalie"
an["B2"].font = f_title
an["B3"] = ("Écart de chaque magasin face au reste du réseau. Ratio = valeur du magasin / valeur des 3 autres magasins le même mois. "
            "Écart = ratio du mois / moyenne du ratio sur les 11 autres mois - 1. "
            "Cette méthode neutralise la saisonnalité commune à tout le réseau.")
an["B3"].font, an["B3"].alignment = f_sub, Alignment(wrap_text=True)
an.merge_cells("B3:H3")
an.row_dimensions[3].height = 28
an["B5"] = "Seuil d'alerte"
an["B5"].font = f_bb
an["C5"] = -0.10
an["C5"].font, an["C5"].fill, an["C5"].number_format, an["C5"].border = f_input, fill_y, PCT, box
an["C5"].comment = Comment("Hypothèse : un écart inférieur à -10 % est considéré comme anormal (le bruit habituel reste sous 5 %).", "Analyse")
nm = len(MAGS)
lastc = L(2 + nm)

def block(top, title, value_key):
    """Bloc valeurs + ratio + écart ; renvoie (ligne_début_écart, ligne_fin_écart)."""
    an.cell(row=top, column=2, value=title).font = f_sec
    header(an, top + 1, 2, ["Mois"] + MAGS + ["Réseau"])
    v0 = top + 2
    for i, m in enumerate(MOIS):
        r = v0 + i
        an.cell(row=r, column=2, value=m)
        for j in range(nm):
            col = L(3 + j)
            an.cell(row=r, column=3 + j, value=f"=SUMIFS({R[value_key]},{R['mois']},$B{r},{R['mag']},{col}${top+1})")
        an.cell(row=r, column=3 + nm, value=f"=SUM(C{r}:{lastc}{r})")
    v1 = v0 + len(MOIS) - 1
    style_range(an, v0, v1, 2, 3 + nm, EUR if value_key == "ca" else INT)

    rt = v1 + 2
    an.cell(row=rt, column=2, value="Ratio magasin / reste du réseau").font = f_bb
    header(an, rt + 1, 2, ["Mois"] + MAGS)
    q0 = rt + 2
    for i, m in enumerate(MOIS):
        r, src = q0 + i, v0 + i
        an.cell(row=r, column=2, value=f"=B{src}")
        for j in range(nm):
            col = L(3 + j)
            an.cell(row=r, column=3 + j, value=f"=IFERROR({col}{src}/(${L(3+nm)}{src}-{col}{src}),0)")
    q1 = q0 + len(MOIS) - 1
    style_range(an, q0, q1, 2, 2 + nm, "0.0000")

    et = q1 + 2
    an.cell(row=et, column=2, value="Écart vs référence (moyenne des 11 autres mois)").font = f_bb
    header(an, et + 1, 2, ["Mois"] + MAGS)
    e0 = et + 2
    for i, m in enumerate(MOIS):
        r, src = e0 + i, q0 + i
        an.cell(row=r, column=2, value=f"=B{src}")
        for j in range(nm):
            col = L(3 + j)
            an.cell(row=r, column=3 + j,
                    value=f"=IFERROR({col}{src}/((SUM({col}${q0}:{col}${q1})-{col}{src})/{len(MOIS)-1})-1,0)")
    e1 = e0 + len(MOIS) - 1
    style_range(an, e0, e1, 2, 2 + nm, PCT)
    rngE = f"C{e0}:{lastc}{e1}"
    an.conditional_formatting.add(rngE, CellIsRule(operator="lessThanOrEqual", formula=["$C$5"],
        fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(name=F, bold=True, color="C00000")))
    return e0, e1

# Synthèse en haut, blocs en dessous
S = 7
an.cell(row=S, column=2, value="Synthèse").font = f_sec
TOP_CA = 16
e0, e1 = block(TOP_CA, "1. Chiffre d'affaires HT", "ca")
TOP_TX = e1 + 3
t0, t1 = block(TOP_TX, "2. Nombre de transactions (fréquentation)", "tx")
RC = f"C{e0}:{lastc}{e1}"
RT = f"C{t0}:{lastc}{t1}"

# Ligne d'aide : pire écart par magasin (pour localiser le min global)
hr = t1 + 3
an.cell(row=hr, column=2, value="Aide au calcul : pire écart CA par magasin").font = f_bb
header(an, hr + 1, 2, ["Indicateur"] + MAGS)
an.cell(row=hr + 2, column=2, value="Écart minimum")
an.cell(row=hr + 3, column=2, value="Mois du minimum")
for j in range(nm):
    col = L(3 + j)
    an.cell(row=hr + 2, column=3 + j, value=f"=MIN({col}{e0}:{col}{e1})").number_format = PCT
    an.cell(row=hr + 3, column=3 + j, value=f"=INDEX($B${e0}:$B${e1},MATCH({col}{hr+2},{col}{e0}:{col}{e1},0))")
style_range(an, hr + 2, hr + 3, 2, 2 + nm)
for j in range(nm):
    an.cell(row=hr + 2, column=3 + j).number_format = PCT

syn = [
    ("Nombre d'alertes CA", f'=COUNTIF({RC},"<="&$C$5)', INT),
    ("Pire écart CA", f"=MIN({RC})", PCT),
    ("Magasin concerné", f"=INDEX($C${hr+1}:${lastc}${hr+1},MATCH(C9,$C${hr+2}:${lastc}${hr+2},0))", None),
    ("Mois concerné", f"=INDEX($C${hr+3}:${lastc}${hr+3},MATCH(C9,$C${hr+2}:${lastc}${hr+2},0))", None),
    ("Écart transactions (même cellule)",
     f"=INDEX({RT},MATCH(C11,$B${t0}:$B${t1},0),MATCH(C10,$C${TOP_TX+1}:${lastc}${TOP_TX+1},0))", PCT),
    ("Écart panier moyen (déduit)", "=IFERROR((1+C9)/(1+C12)-1,0)", PCT),
]
for i, (lab, fml, fmt) in enumerate(syn):
    r = S + 1 + i
    an.cell(row=r, column=2, value=lab).font = f_b
    c = an.cell(row=r, column=3, value=fml)
    c.font, c.border = f_bb, box
    if fmt:
        c.number_format = fmt
an["D9"] = '=IF(C9<=$C$5,"ANOMALIE DÉTECTÉE","Aucune anomalie")'
an["D9"].font = Font(name=F, bold=True, color="C00000")
an["B14"] = ('=IF(C8=0,"Aucun écart ne dépasse le seuil.",'
             '"Diagnostic : "&C10&" en "&C11&" affiche un CA "&TEXT(-C9,"0.0%")&" sous sa référence. '
             'Les transactions baissent de "&TEXT(-C12,"0.0%")&" alors que le panier moyen varie de "&TEXT(C13,"0.0%")&" : '
             'la baisse vient de la fréquentation, pas du prix ni du mix produit.")')
an["B14"].font, an["B14"].alignment = Font(name=F, size=10, bold=True, color=NAVY), Alignment(wrap_text=True, vertical="top")
an.merge_cells("B14:H14")
an.row_dimensions[14].height = 42

wb.move_sheet("TCD", offset=0)
out = ROOT / "excel" / "analyse_ventes_retail_2025.xlsx"
wb.save(out)
print("Classeur écrit :", out)
